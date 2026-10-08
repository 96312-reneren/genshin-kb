#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
把项目根目录的 Markdown 报告转换为可编辑的 .docx。

用法:
    python scripts/md_to_docx.py 报告.md [输出.docx]

只依赖 python-docx，不联网。转换规则覆盖本项目报告用到的全部语法:
标题(#~####)、表格、引用块、无序/有序列表、围栏代码块、水平线、行内 **粗体** / *斜体* / `代码`。
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

CJK_FONT = "微软雅黑"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
MUTED = RGBColor(0x59, 0x59, 0x59)


# ---------------------------------------------------------------- 底层工具

def set_run_font(run, name=CJK_FONT, size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(element, fill):
    """给段落或单元格加底色。接受 Paragraph / _Cell / lxml 元素。"""
    node = getattr(element, "_p", None)
    if node is None:
        node = getattr(element, "_tc", None)
    if node is None:
        node = element
    pr = node.get_or_add_tcPr() if node.tag.endswith("}tc") else node.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pr.append(shd)


def para_border(paragraph, edge="left", size=18, color="1F4E79"):
    ppr = paragraph._p.get_or_add_pPr()
    borders = ppr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        ppr.append(borders)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), "8")
    el.set(qn("w:color"), color)
    borders.append(el)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*]+\*)")


def add_inline(paragraph, text, base_size=10.5):
    """写入带行内格式的文本。"""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = paragraph.add_run(part[2:-2])
            set_run_font(r, size=base_size, bold=True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = paragraph.add_run(part[1:-1])
            set_run_font(r, name=MONO_FONT, size=base_size - 0.5, color=RGBColor(0xC0, 0x39, 0x2B))
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = paragraph.add_run(part[1:-1])
            set_run_font(r, size=base_size, italic=True)
        else:
            r = paragraph.add_run(part)
            set_run_font(r, size=base_size)


# ---------------------------------------------------------------- 构建文档

def build_doc(md_text, out_path):
    doc = Document()

    # 全局正文样式
    normal = doc.styles["Normal"]
    normal.font.name = CJK_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.4

    for i, sz in enumerate([22, 16, 13, 11.5], start=1):
        st = doc.styles[f"Heading {i}"]
        st.font.name = CJK_FONT
        st.font.size = Pt(sz)
        st.font.color.rgb = ACCENT
        st.font.bold = True
        st.element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)

    lines = md_text.replace("\r\n", "\n").split("\n")
    i, n = 0, len(lines)
    first_h1_done = False

    while i < n:
        line = lines[i]
        stripped = line.strip()

        # 围栏代码块
        if stripped.startswith("```"):
            i += 1
            buf = []
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            for code_line in buf:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.left_indent = Pt(12)
                shade(p, "F2F2F2")
                r = p.add_run(code_line if code_line.strip() else " ")
                set_run_font(r, name=MONO_FONT, size=9.5)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue

        # 表格
        if stripped.startswith("|") and i + 1 < n and set(lines[i + 1].replace("|", "").strip()) <= {"-", ":", " "} and "|" in lines[i + 1]:
            header = [c.strip() for c in stripped.strip("|").split("|")]
            rows = []
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for ci, htxt in enumerate(header):
                cell = table.rows[0].cells[ci]
                cell.text = ""
                shade(cell._tc, "1F4E79")
                add_inline(cell.paragraphs[0], htxt, base_size=9.5)
                for r in cell.paragraphs[0].runs:
                    r.font.bold = True
                    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            for ri, row in enumerate(rows):
                cells = table.add_row().cells
                for ci in range(len(header)):
                    val = row[ci] if ci < len(row) else ""
                    cells[ci].text = ""
                    add_inline(cells[ci].paragraphs[0], val, base_size=9.5)
                    if ri % 2 == 1:
                        shade(cells[ci]._tc, "F7F9FC")
            doc.add_paragraph().paragraph_format.space_after = Pt(4)
            continue

        # 水平线
        if stripped in {"---", "***", "___"}:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            para_border(p, edge="bottom", size=6, color="BFBFBF")
            i += 1
            continue

        # 引用块
        if stripped.startswith(">"):
            buf = []
            while i < n and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            for q in buf:
                if not q:
                    continue
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(14)
                p.paragraph_format.space_after = Pt(4)
                para_border(p, edge="left", size=18, color="1F4E79")
                add_inline(p, q, base_size=10)
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level, text = len(m.group(1)), m.group(2)
            if level == 1 and not first_h1_done:
                first_h1_done = True
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(14)
                add_inline(p, text, base_size=22)
                for r in p.runs:
                    r.font.bold = True
                    r.font.color.rgb = ACCENT
            else:
                doc.add_heading(text, level=min(level, 4))
            i += 1
            continue

        # 无序列表
        m = re.match(r"^[-*]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, m.group(1))
            i += 1
            continue

        # 有序列表
        m = re.match(r"^\d+[.)]\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            add_inline(p, m.group(1))
            i += 1
            continue

        # 空行
        if not stripped:
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        add_inline(p, stripped)
        i += 1

    doc.save(out_path)
    return out_path


def set(x):
    """表格分隔行判定：收集去掉 | 后的字符集合。"""
    return {c for c in x if c}


def main():
    if len(sys.argv) < 2:
        print("用法: python scripts/md_to_docx.py <输入.md> [输出.docx]")
        return 1
    src = Path(sys.argv[1])
    if not src.exists():
        print(f"找不到输入文件: {src}")
        return 1
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src.with_suffix(".docx")
    build_doc(src.read_text(encoding="utf-8"), dst)
    print(f"已生成: {dst}  ({dst.stat().st_size} 字节)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
